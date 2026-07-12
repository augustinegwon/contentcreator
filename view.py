#!/usr/bin/env python3
"""
수집한 데이터를 엑셀 파일(.xlsx) 하나로 뽑아주는 도구.

사용법:
    python view.py

실행하면 같은 폴더에 youtube_data.xlsx 파일 1개가 만들어진다.
그 파일 안에 시트(탭) 3개가 들어 있다:
    - 수집기록 : 언제 무슨 키워드를 수집했는지 (수집 이벤트 목록)
    - 영상     : 수집한 영상들 (제목, 조회수, 좋아요 등)
    - 댓글     : 수집한 댓글들 (작성자, 내용, 좋아요 등)

이 파일을 더블클릭하면 엑셀(Excel)이나 넘버스(Numbers)로 바로 열린다.
"""

import os
import sqlite3
import sys

# openpyxl: 파이썬에서 엑셀 .xlsx 파일을 만드는 라이브러리
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

DB_FILE = "youtube_data.db"          # collect.py 가 만든 데이터 파일
OUTPUT_FILE = "youtube_data.xlsx"    # 만들어질 엑셀 파일


# (시트이름, SQL 쿼리) 목록. 시트 하나당 표 하나가 들어간다.
# 사람이 보기 좋은 순서로 컬럼을 정렬해서 가져온다.
SHEETS = {
    "수집기록": """
        SELECT id AS 수집번호,
               keyword AS 키워드,
               collected_at AS 수집시각,
               video_count AS 영상수,
               comment_count AS 댓글수
        FROM searches
        ORDER BY id;
    """,
    "영상": """
        SELECT search_id AS 수집번호,
               title AS 제목,
               channel_title AS 채널,
               view_count AS 조회수,
               like_count AS 좋아요,
               comment_count AS 댓글수,
               published_at AS 게시일,
               url AS 링크,
               collected_at AS 수집시각
        FROM videos
        ORDER BY search_id, view_count DESC;
    """,
    "댓글": """
        SELECT search_id AS 수집번호,
               video_id AS 영상ID,
               author AS 작성자,
               text AS 댓글내용,
               like_count AS 좋아요,
               published_at AS 작성일,
               collected_at AS 수집시각
        FROM comments
        ORDER BY search_id, like_count DESC;
    """,
}


def main() -> None:
    # 1) 데이터 파일이 있는지 먼저 확인한다.
    if not os.path.exists(DB_FILE):
        print(
            f"[에러] '{DB_FILE}' 파일이 없습니다.\n"
            "       먼저 수집을 한 번 해야 해요:  python collect.py \"키워드\"",
            file=sys.stderr,
        )
        sys.exit(1)

    conn = sqlite3.connect(DB_FILE)

    # 2) 엑셀 파일(워크북)을 새로 만든다.
    wb = Workbook()
    wb.remove(wb.active)  # openpyxl 이 기본으로 넣는 빈 시트를 지운다.

    header_font = Font(bold=True)  # 첫 줄(제목 줄)은 굵게

    total_rows = 0
    for sheet_name, query in SHEETS.items():
        cursor = conn.execute(query)
        headers = [d[0] for d in cursor.description]  # 컬럼 이름(한글)
        rows = cursor.fetchall()

        ws = wb.create_sheet(title=sheet_name)

        # 첫 줄: 컬럼 제목 (굵게)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = header_font

        # 나머지 줄: 실제 데이터
        for row in rows:
            ws.append(list(row))

        # 보기 좋게: 첫 줄 고정(스크롤해도 제목이 남음) + 컬럼 너비 자동 조정
        ws.freeze_panes = "A2"
        _auto_width(ws, headers, rows)

        print(f"  [{sheet_name}] 시트 → {len(rows)}줄")
        total_rows += len(rows)

    conn.close()

    # 3) 저장
    wb.save(OUTPUT_FILE)

    path = os.path.abspath(OUTPUT_FILE)
    print("\n" + "=" * 60)
    print(f"엑셀 파일 저장 완료 (시트 3개, 총 {total_rows}줄)")
    print(f"위치: {path}")
    print("이 파일을 더블클릭하면 엑셀/넘버스로 열려요. (아래 탭으로 시트 전환)")
    print("=" * 60)


def _auto_width(ws, headers, rows) -> None:
    """각 컬럼 너비를 내용 길이에 맞춰 대충 조정한다. (너무 넓어지지 않게 상한선 둠)"""
    for col_idx in range(len(headers)):
        # 그 컬럼에서 가장 긴 글자 수를 찾는다. (제목 + 데이터)
        longest = len(str(headers[col_idx]))
        for row in rows:
            value = row[col_idx]
            if value is not None:
                longest = max(longest, len(str(value)))
        # 너비 = 글자 수 + 여유 2칸, 단 최대 60칸까지만
        width = min(longest + 2, 60)
        ws.column_dimensions[get_column_letter(col_idx + 1)].width = width


if __name__ == "__main__":
    main()
